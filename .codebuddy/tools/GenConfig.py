import openpyxl, json, os

SRC = "addons/TinyTools/ConfigGenerator/"
OUT = "Resources/Config/"

def to_str(v):
    if v is None: return ""
    if isinstance(v, bool): return "True" if v else "False"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)

def split_arr(s):
    return [p.strip() for p in s.split("|") if p.strip() != ""]

def parse_bool(s):
    s = s.lower()
    return s in ("true", "1", "yes")

def convert(t, v):
    s = to_str(v).strip()
    if s == "": return None
    if t == "int":
        try: return int(s)
        except: return 0
    if t == "long":
        try: return int(s)
        except: return 0
    if t == "float":
        try:
            f = float(s); return int(f) if f.is_integer() else f
        except: return 0.0
    if t == "double":
        try:
            f = float(s); return int(f) if f.is_integer() else f
        except: return 0.0
    if t == "bool":
        return parse_bool(s)
    if t == "string":
        return s
    if t == "int[]":
        return [int(x) for x in split_arr(s)]
    if t == "long[]":
        return [int(x) for x in split_arr(s)]
    if t == "float[]":
        return [float(x) for x in split_arr(s)]
    if t == "double[]":
        return [float(x) for x in split_arr(s)]
    if t == "bool[]":
        return [parse_bool(x) for x in split_arr(s)]
    if t == "string[]":
        return split_arr(s)
    # default: 枚举 / BigNum -> 能转 int 就转, 否则保留原始字符串(与 C# default 分支一致)
    try: return int(s)
    except: return s

def gen_one(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    name = ws.title
    maxc = ws.max_column
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, maxc + 1)])
    names = rows[0]; types = rows[1]
    out = []
    for row in rows[4:]:
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        d = {}
        for j in range(len(names)):
            if j >= len(row): break
            field = names[j]
            if not field: continue
            val = row[j]
            if val is None: continue
            cv = convert(types[j], val)
            if cv is None: continue
            d[field] = cv
        out.append(d)
    os.makedirs(OUT, exist_ok=True)
    with open(OUT + name + ".json", "w", encoding="utf-8", newline="\n") as f:
        json.dump(out, f, indent=2, ensure_ascii=True)
    print(f"generated {name}.json ({len(out)} rows)")

if __name__ == "__main__":
    for fn in os.listdir(SRC):
        if fn.lower().endswith(".xlsx"):
            gen_one(os.path.join(SRC, fn))
    print("ALL DONE")
